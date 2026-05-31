"""Сборка docx с двумя примерами выдач (baseline vs ours) для нейтральных
статей-источников. Используется тот же production-пайплайн, что и в
make_neutral_examples.py, но без столбца "департамент" и с аккуратной
вёрсткой под вставку в диплом.

Выход: reports/examples_neutral/examples_neutral.docx
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Pt

from make_neutral_examples import (
    K,
    baseline_topk,
    find_neutral_sources,
    our_pipeline,
)
from serve_carousel import build_context
from train_ltr_full import REPORTS_DIR, safe_str


EXAMPLES_DIR = REPORTS_DIR / "examples_neutral"
EXAMPLES_DIR.mkdir(exist_ok=True)
OUT_PATH = EXAMPLES_DIR / "examples_neutral.docx"


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    run = p.add_run(text)
    run.font.size = Pt(size)
    if bold:
        run.bold = True
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _add_table(doc: Document, df: pd.DataFrame, *, caption: str) -> None:
    cap = doc.add_paragraph()
    cap_run = cap.add_run(caption)
    cap_run.bold = True
    cap_run.font.size = Pt(10)

    cols = list(df.columns)
    table = doc.add_table(rows=len(df) + 1, cols=len(cols))
    table.style = "Light Grid Accent 1"
    table.autofit = True

    for j, col in enumerate(cols):
        _set_cell_text(table.cell(0, j), col, bold=True)

    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(cols):
            val = row[col]
            if isinstance(val, float):
                txt = f"{val:.3f}"
            else:
                txt = str(val)
            _set_cell_text(table.cell(i, j), txt)

    doc.add_paragraph("")


def _prepare(df: pd.DataFrame, columns_order: list[str]) -> pd.DataFrame:
    """Уберём столбец department и колонку rubric, оставим заголовок,
    свежесть/просмотры (для baseline) и метрики ранжирования."""
    cols = [c for c in columns_order if c in df.columns]
    out = df[cols].copy()
    return out.reset_index(drop=True)


def main() -> None:
    print("[serve] build context…")
    ctx = build_context()
    art = ctx.art

    sources = find_neutral_sources(ctx)
    if not sources:
        print("не нашли подходящих источников")
        return
    print("sources:", sources)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    # Альбомная ориентация и узкие поля — чтобы длинные заголовки помещались целиком.
    for s in doc.sections:
        s.orientation = WD_ORIENT.LANDSCAPE
        new_width, new_height = s.page_height, s.page_width
        s.page_width = new_width
        s.page_height = new_height
        s.left_margin = Cm(1.2)
        s.right_margin = Cm(1.2)
        s.top_margin = Cm(1.2)
        s.bottom_margin = Cm(1.2)

    h = doc.add_paragraph()
    r = h.add_run("Пример выдач: baseline (cosine kNN) vs наш пайплайн (similar + explore → LTR → rerank, интерливинг 4/7/10)")
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph(
        "Для каждого из шести нейтральных источников приводится top-12 baseline-ранжирования "
        "(только cosine similarity по текстовым эмбеддингам, без LTR и продуктовых правил) "
        "и финальная карусель нашего пайплайна. У нашего пайплайна 9 слотов отведены пулу "
        "similar (kNN ∪ coview), 3 слота — пулу explore с интерливингом на позициях 4, 7 и 10. "
        "Заголовки приведены полностью, без обрезки. Пятый и шестой источники специально выбраны "
        "так, чтобы софт-макс семплирование explore-пула выдало другие материалы — это иллюстрирует, "
        "что разнообразие между источниками действительно работает."
    )

    for i, sid in enumerate(sources, start=1):
        src_idx = art.article_id_to_row.get(sid)
        if src_idx is None:
            continue
        src_title = safe_str(art.df.iloc[src_idx].get("article_base__title", ""))
        src_rub = safe_str(art.df.iloc[src_idx].get("article_base__rubric", ""))

        p = doc.add_paragraph()
        r = p.add_run(f"Источник №{i}: «{src_title}»")
        r.bold = True
        r.font.size = Pt(11)
        doc.add_paragraph(f"Рубрика: {src_rub}. article_id: {sid}.")

        base_df = baseline_topk(ctx, src_idx, k=K)
        ours_df = our_pipeline(ctx, sid)

        base_show = _prepare(
            base_df,
            ["rank", "title", "age_days", "views", "sim"],
        )
        ours_show = _prepare(
            ours_df,
            ["rank", "mix", "title", "sim", "ltr_score"],
        )

        _add_table(doc, base_show, caption=f"Таблица — baseline (cosine kNN top-{K}) для источника №{i}")
        _add_table(doc, ours_show, caption=f"Таблица — наш пайплайн (LTR + rerank, K={K}) для источника №{i}")

        doc.add_paragraph()

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
